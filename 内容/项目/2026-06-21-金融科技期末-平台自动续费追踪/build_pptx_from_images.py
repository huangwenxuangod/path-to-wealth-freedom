"""14 张 PNG → 1 个 14 页 pptx
每页 16:9 (1920x1080 → 13.33 x 7.5 inch)
保留 16:9 高清版（每页 1 张图占满 slide）
"""
import os
from pptx import Presentation
from pptx.util import Inches, Emu

SRC_DIR = r"D:\path-to-wealth-freedom\.workbuddy\deck-images"
OUT = r"D:\path-to-wealth-freedom\内容\项目\2026-06-21-金融科技期末-平台自动续费追踪\金融科技应用前沿-醒账-答辩PPT.pptx"

# 16:9 slide size: 13.333 x 7.5 inch
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]  # Blank

for i in range(1, 15):
    img = os.path.join(SRC_DIR, f"slide-{i:02d}.png")
    if not os.path.exists(img):
        print(f"missing {img}")
        continue
    slide = prs.slides.add_slide(blank_layout)
    # 全屏占满 slide
    slide.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
    print(f"  page {i}: {os.path.basename(img)}")

prs.save(OUT)
size = os.path.getsize(OUT) / 1024
print(f"\nDone: {OUT}")
print(f"Pages: {len(prs.slides)}")
print(f"Size: {size:.0f} KB")
